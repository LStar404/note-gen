"""PPTX 讲义解析（5.1 扩展）：幻灯片 = 页，文本框 → 块，插图 → figures/。

老师按份发放的 PPT 讲义直接导入：合成与 PDF 解析完全同形的 pages JSON 与 manifest，
锚点/生成/检索整条管线原样复用（python-pptx 内置抽取，无需外部转换器）。
PDF 面板暂无预览（本次只做解析适配，不做导出适配）。
"""

from __future__ import annotations

import json
from pathlib import Path

from src.core.jsonl import atomic_write_bytes, canonical_json, sha256_file
from src.core.models import Manifest
from src.core.paths import check_doc_id, rel_to_root
from src.core.text import utc_now_iso
from src.ingest.chapter_detect import detect_chapters

EMU_PER_PT = 12700.0


def _bbox_from_shape(shape, h_pt: float) -> list[float] | None:
    """pptx 顶左原点（EMU）→ PDF 用户空间底左（点）；无几何的形状不产出块。"""
    try:
        x0 = float(shape.left) / EMU_PER_PT
        y0 = float(shape.top) / EMU_PER_PT
        x1 = x0 + float(shape.width) / EMU_PER_PT
        y1 = y0 + float(shape.height) / EMU_PER_PT
    except Exception:
        return None
    if x1 <= x0 or y1 <= y0:
        return None
    return [x0, h_pt - y1, x1, h_pt - y0]


def _is_title(shape) -> bool:
    try:
        return bool(shape.is_placeholder) and int(shape.placeholder_format.type) in (1, 3)
    except Exception:
        return False


def _shape_blocks(shape, ctx: dict) -> list[dict]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out: list[dict] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            out.extend(_shape_blocks(sub, ctx))
        return out
    bbox = _bbox_from_shape(shape, ctx["h_pt"])
    if bbox is None:
        return out
    if getattr(shape, "has_text_frame", False):
        text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
        if text:
            out.append({"type": "heading" if _is_title(shape) else "paragraph",
                        "level": 2 if _is_title(shape) else 0,
                        "text": text, "bbox": bbox})
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            ctx["figure_seq"] += 1
            fig_id = f"fig_{ctx['figure_seq']:08d}"
            fname = f"{fig_id}.png"
            (ctx["figures_dir"] / fname).write_bytes(shape.image.blob)
            out.append({"type": "figure", "level": 0, "text": "（幻灯片插图）",
                        "bbox": bbox, "figure_id": fig_id,
                        "image_path": rel_to_root(ctx["root"], ctx["figures_dir"] / fname)})
        except Exception:
            pass
    return out


def parse_pptx(pptx_path: str | Path, out_dir: str | Path, *, layout, doc_id: str | None = None,
               profile: str = "textbook", title: str | None = None,
               chapters_override: list | None = None,
               progress=None) -> dict:
    """解析 PPTX → pages/*.json + manifest（与 parse_pdf 同契约）。"""
    from pptx import Presentation

    pptx_path = Path(pptx_path).resolve()
    if not pptx_path.exists():
        raise FileNotFoundError(f"PPTX 不存在: {pptx_path}")
    out_dir = Path(out_dir)
    doc_id = doc_id or ("doc-" + sha256_file(pptx_path)[:12])
    check_doc_id(doc_id)

    prs = Presentation(str(pptx_path))
    w_pt = float(prs.slide_width) / EMU_PER_PT
    h_pt = float(prs.slide_height) / EMU_PER_PT
    view_box = [0.0, 0.0, w_pt, h_pt]
    n_slides = len(prs.slides)
    if n_slides == 0:
        raise ValueError("PPTX 没有任何幻灯片")

    pages_dir = out_dir / doc_id / "pages"
    doc_dir = layout.doc_dir(doc_id)
    figures_dir = layout.figures_dir(doc_id)
    for sub in ("anchors", "blocks", "history", "figures", "refs"):
        (doc_dir / sub).mkdir(parents=True, exist_ok=True)
    if pages_dir.exists():
        import shutil

        shutil.rmtree(pages_dir)
    pages_dir.mkdir(parents=True, exist_ok=True)

    ctx = {"h_pt": h_pt, "root": layout.root, "figures_dir": figures_dir, "figure_seq": 0}
    page_records: list[dict] = []
    failed = 0
    for i, slide in enumerate(prs.slides):
        if progress and (i % 5 == 0 or i == n_slides - 1):
            progress({"stage": "ingest", "pages_done": i, "pages_total": n_slides,
                      "parser": "pptx"})
        blocks: list[dict] = []
        try:
            for shape in slide.shapes:
                blocks.extend(_shape_blocks(shape, ctx))
        except Exception:
            failed += 1
        status = "ok" if blocks else "blank"
        page_json = {
            "doc_id": doc_id,
            "page": i + 1,
            "page_width": w_pt,
            "page_height": h_pt,
            "coordinate_meta": {
                "source_space": "pptx_user_space",
                "source_to_pdf": [1.0, 0.0, 0.0, 1.0, 0.0, 0.0],
                "view_box": view_box,
                "user_unit": 1.0,
                "rotation": 0,
            },
            "parse_status": status,
            "blocks": blocks,
        }
        atomic_write_bytes(pages_dir / f"{i + 1:04d}.json",
                           (json.dumps(page_json, ensure_ascii=False, indent=2) + "\n").encode("utf-8"),
                           fsync=False)
        page_records.append({"page": i + 1, "view_box": list(view_box),
                             "user_unit": 1.0, "rotation": 0, "parse_status": status})
    if n_slides and failed / n_slides > 0.05:
        raise RuntimeError(f"失败页数 {failed}/{n_slides} 超过 5%")

    if progress:
        progress({"stage": "ingest", "pages_done": n_slides, "pages_total": n_slides,
                  "parser": "pptx"})
    chapters = detect_chapters(None, page_records, doc_id, pages_dir=pages_dir,
                               profile=profile, override=chapters_override)

    from src.ingest.parse_pdf import PAGE_JSON_VERSION, _pages_hash

    now = utc_now_iso()
    manifest = {
        "schema_version": 2,
        "doc_id": doc_id,
        "profile": profile,
        "title": title or pptx_path.stem,
        "pdf_path": rel_to_root(layout.root, pptx_path),   # 无 PDF：指向原始 PPTX
        "pdf_sha256": sha256_file(pptx_path),
        "parser": {"name": "pptx", "version": str(PAGE_JSON_VERSION),
                   "config_hash": "sha256:" + __import__("hashlib").sha256(canonical_json(
                       {"name": "pptx", "profile": profile}).encode("utf-8")).hexdigest()},
        "ocr": False,
        "data_revision": 1,
        "pages": page_records,
        "collision_map": {},
        "chapters": chapters,
        "stage_hashes": {
            "ingest": {
                "input_hash": "sha256:" + sha256_file(pptx_path),
                "output_hash": "sha256:" + _pages_hash(pages_dir),
                "version": PAGE_JSON_VERSION,
            }
        },
        "created_at": now,
        "updated_at": now,
    }
    Manifest.model_validate(manifest)
    atomic_write_bytes(layout.manifest(doc_id),
                       (json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    return manifest
