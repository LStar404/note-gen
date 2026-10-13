"""PPTX 解析与讲义追加回归：幻灯片合成同形产物；追加后已有笔记必须原样保留。"""

from __future__ import annotations

from fastapi.testclient import TestClient

from src.api.main import create_app
from src.core.store import DocStore


def _tiny_pptx(path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "第 1 章 绪论"
    slide.placeholders[1].text = "本章讨论集合与映射的基本概念。"
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "第 2 章 群作用"
    slide2.placeholders[1].text = "群在集合上的作用与轨道计数。"
    prs.save(str(path))
    return path


def test_parse_pptx_produces_pipeline_compatible_output(tmp_path, workspace):
    """幻灯片 = 页：pages JSON/manifest 与 PDF 解析同契约，锚定管线可直接消费。"""
    from src.anchor.chunker import chunk_and_anchor
    from src.ingest.parse_pptx import parse_pptx

    pptx = _tiny_pptx(tmp_path / "lecture.pptx")
    man = parse_pptx(pptx, workspace.ingest_root, layout=workspace, doc_id="pptx-la",
                     title="讲义")
    assert man["chapters"], man
    assert len(man["pages"]) == 2
    assert man["pdf_path"].endswith("lecture.pptx")
    total = sum(chunk_and_anchor("pptx-la", c["chapter_id"], layout=workspace)
                for c in man["chapters"])
    assert total >= 2, f"标题页/内容页都要出锚点: {total}"


def test_append_lecture_preserves_existing_notes(sample_doc):
    """追加讲义 = 合并后重解析重锚定：已有笔记必须原样保留，新页只带来新章节。"""
    layout, doc_id = sample_doc
    store = DocStore(layout, doc_id)
    ch = store.chapter_ids()[0]
    block = store.create_block({"type": "motivation", "title": "我的手写心得",
                                "content": "这是人工写的重要心得。",
                                "sources": [], "basis": "synthesis", "order": 1.0},
                               author="human", chapter_id=ch)
    before_pages = len(store.manifest()["pages"])
    # 追加的前提：教材源文件在 textbooks/<doc>.pdf（真实导入流会放好；夹具补上）
    import shutil

    layout.textbooks_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy(layout.root / "sample.pdf", layout.textbooks_dir / f"{doc_id}.pdf")

    app = create_app(layout.root)
    pdf_bytes = (layout.root / "sample.pdf").read_bytes()
    with TestClient(app, base_url="http://127.0.0.1:8787") as client:
        headers = {"X-Local-Token": app.state.app_state.token_value}
        import io

        resp = client.post(f"/api/v1/docs/{doc_id}/append",
                           files={"file": ("extra.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
                           headers=headers)
        assert resp.status_code == 202, resp.text
        task_id = resp.json()["task_id"]
        import time

        for _ in range(400):
            t = client.get(f"/api/v1/tasks/{task_id}", headers=headers).json()
            if t["status"] not in ("queued", "running"):
                break
            time.sleep(0.15)
        assert t["status"] == "succeeded", t
    store2 = DocStore(layout, doc_id)
    after_pages = len(store2.manifest()["pages"])
    assert after_pages > before_pages, f"页数必须增加: {before_pages} -> {after_pages}"
    kept = [b for b in store2.block_rows(ch) if b["id"] == block["id"]]
    assert kept and kept[0]["content"] == "这是人工写的重要心得。", "追加不得丢已有笔记"


def test_import_rejects_unsupported_format(sample_doc):
    layout, doc_id = sample_doc
    app = create_app(layout.root)
    with TestClient(app, base_url="http://127.0.0.1:8787") as client:
        headers = {"X-Local-Token": app.state.app_state.token_value}
        import io

        resp = client.post("/api/v1/docs",
                           files={"file": ("notes.docx", io.BytesIO(b"xx"), "application/msword")},
                           headers=headers)
        assert resp.status_code >= 400, resp.text
        assert "PDF" in resp.text or "PPTX" in resp.text
